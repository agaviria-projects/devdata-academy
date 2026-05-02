import streamlit as st

st.set_page_config(page_title="Python", page_icon="🐍", layout="wide")

st.title("🐍 Python")
st.subheader("Bases, comandos y lógica para automatización, datos y desarrollo")

st.markdown("""
## 📘 ¿Qué es Python?

Python es un lenguaje de programación que permite crear scripts, automatizar tareas, limpiar datos, procesar archivos, conectarse a bases de datos y construir aplicaciones.

En tus proyectos lo has usado para:

- Limpiar archivos Excel
- Procesar datos de inventario
- Automatizar PDFs
- Crear apps con Streamlit
- Trabajar con SQLite
- Preparar información para Power BI
""")

st.divider()

st.markdown("## 🧠 Forma correcta de pensar en Python")

st.info("""
No se trata de copiar y pegar código.

La lógica correcta es:

1. ¿Qué problema tengo?
2. ¿Qué datos necesito?
3. ¿Dónde están esos datos?
4. ¿Qué transformación debo hacer?
5. ¿Qué resultado espero?
6. ¿Dónde guardo la salida?
""")

st.code("""
# Ejemplo mental
# Problema: tengo un Excel con ventas y necesito calcular total.
# Datos: cantidad y precio.
# Transformación: total = cantidad * precio.
# Salida: nuevo Excel limpio.
""", language="python")

st.divider()

st.markdown("## 🧩 Estructura básica de un script")

st.code("""
# 1. Importar librerías
import pandas as pd
from pathlib import Path

# 2. Definir rutas
ruta_entrada = Path("data/ventas.xlsx")
ruta_salida = Path("output/ventas_limpias.xlsx")

# 3. Cargar datos
df = pd.read_excel(ruta_entrada)

# 4. Procesar datos
df["total"] = df["cantidad"] * df["precio"]

# 5. Exportar resultado
df.to_excel(ruta_salida, index=False)

print("Proceso finalizado correctamente")
""", language="python")

st.markdown("""
### ¿Qué hace cada parte?

- `import`: trae herramientas externas.
- `Path`: ayuda a manejar rutas.
- `pd.read_excel`: lee un archivo Excel.
- `df["total"]`: crea una nueva columna.
- `to_excel`: exporta el resultado.
- `print`: muestra mensajes en consola.
""")

st.divider()

st.markdown("## 📦 Librerías que más usas")

st.markdown("""
### pandas
Sirve para trabajar con tablas de datos.

### openpyxl
Sirve para leer, escribir y formatear archivos Excel.

### pathlib
Sirve para manejar rutas de carpetas y archivos.

### os
Sirve para interactuar con el sistema operativo.

### sqlite3
Sirve para conectarse a bases de datos SQLite.

### streamlit
Sirve para crear aplicaciones web con Python.

### fitz / PyMuPDF
Sirve para trabajar con PDF.

### PIL / Pillow
Sirve para trabajar con imágenes.
""")

st.divider()

st.markdown("## 💻 Comandos esenciales de Python")

st.markdown("### Variables")

st.code("""
nombre = "Héctor"
edad = 30
activo = True
""", language="python")

st.info("Una variable guarda un valor para usarlo después.")

st.markdown("### Condicionales")

st.code("""
stock = 5

if stock > 0:
    print("Hay inventario")
else:
    print("No hay inventario")
""", language="python")

st.info("Sirven para tomar decisiones.")

st.markdown("### Bucles")

st.code("""
productos = ["Router", "ONT", "Cable"]

for producto in productos:
    print(producto)
""", language="python")

st.info("Sirven para repetir una acción sobre varios elementos.")

st.markdown("### Funciones")

st.code("""
def calcular_total(cantidad, precio):
    total = cantidad * precio
    return total

resultado = calcular_total(5, 20000)
print(resultado)
""", language="python")

st.info("Una función guarda lógica reutilizable.")

st.divider()

st.markdown("## 🐼 pandas: lo más importante")

st.markdown("### Leer archivos")

st.code("""
import pandas as pd

df = pd.read_excel("ventas.xlsx")
df_csv = pd.read_csv("ventas.csv")
""", language="python")

st.markdown("### Ver datos")

st.code("""
df.head()
df.tail()
df.info()
df.describe()
df.columns
""", language="python")

st.markdown("""
### ¿Para qué sirve?

- `head()`: ver primeras filas.
- `tail()`: ver últimas filas.
- `info()`: ver columnas y tipos de datos.
- `describe()`: ver estadísticas.
- `columns`: ver nombres de columnas.
""")

st.markdown("### Seleccionar columnas")

st.code("""
df["cliente"]

df[["cliente", "ciudad", "total"]]
""", language="python")

st.markdown("### Crear columnas")

st.code("""
df["total"] = df["cantidad"] * df["precio"]
""", language="python")

st.markdown("### Filtrar datos")

st.code("""
df_medellin = df[df["ciudad"] == "Medellín"]

df_altas = df[df["total"] > 100000]
""", language="python")

st.markdown("### Agrupar datos")

st.code("""
ventas_por_ciudad = df.groupby("ciudad")["total"].sum().reset_index()
""", language="python")

st.markdown("### Ordenar datos")

st.code("""
df_ordenado = df.sort_values("total", ascending=False)
""", language="python")

st.divider()

st.markdown("## 🧹 Limpieza de datos")

st.markdown("### Ver valores nulos")

st.code("""
df.isna().sum()
""", language="python")

st.markdown("### Rellenar nulos")

st.code("""
df["cantidad"] = df["cantidad"].fillna(0)
df["cliente"] = df["cliente"].fillna("Sin cliente")
""", language="python")

st.markdown("### Eliminar duplicados")

st.code("""
df = df.drop_duplicates()
""", language="python")

st.markdown("### Limpiar texto")

st.code("""
df["cliente"] = df["cliente"].astype(str).str.strip().str.upper()
""", language="python")

st.markdown("### Convertir números")

st.code("""
df["precio"] = pd.to_numeric(df["precio"], errors="coerce").fillna(0)
""", language="python")

st.markdown("### Convertir fechas")

st.code("""
df["fecha"] = pd.to_datetime(df["fecha"], errors="coerce")
""", language="python")

st.warning("""
`errors="coerce"` convierte valores inválidos en NaN o NaT.
Es útil para detectar datos sucios sin romper el script.
""")

st.divider()

st.markdown("## 📁 Manejo de carpetas y archivos")

st.code("""
from pathlib import Path

carpeta_origen = Path("data")
carpeta_salida = Path("output")

carpeta_salida.mkdir(parents=True, exist_ok=True)

archivos_excel = list(carpeta_origen.glob("*.xlsx"))

for archivo in archivos_excel:
    print(archivo.name)
""", language="python")

st.markdown("""
### ¿Para qué sirve?

- Crear carpetas si no existen.
- Buscar archivos por extensión.
- Procesar varios archivos automáticamente.
""")

st.divider()

st.markdown("## 📘 Excel con Python")

st.markdown("### Exportar a Excel")

st.code("""
df.to_excel("ventas_limpias.xlsx", index=False)
""", language="python")

st.markdown("### Exportar varias hojas")

st.code("""
with pd.ExcelWriter("reporte.xlsx", engine="openpyxl") as writer:
    df.to_excel(writer, sheet_name="Datos", index=False)
    ventas_por_ciudad.to_excel(writer, sheet_name="Resumen", index=False)
""", language="python")

st.markdown("### Crear tabla estructurada con openpyxl")

st.code("""
from openpyxl import load_workbook
from openpyxl.worksheet.table import Table, TableStyleInfo

archivo = "ventas_limpias.xlsx"

wb = load_workbook(archivo)
ws = wb.active

tabla = Table(displayName="TablaVentas", ref=ws.dimensions)

estilo = TableStyleInfo(
    name="TableStyleMedium9",
    showRowStripes=True
)

tabla.tableStyleInfo = estilo
ws.add_table(tabla)

wb.save(archivo)
""", language="python")

st.divider()

st.markdown("## 🗄️ SQLite con Python")

st.code("""
import sqlite3

conn = sqlite3.connect("data/conocimiento.db")
cursor = conn.cursor()

cursor.execute("SELECT * FROM conocimiento")

datos = cursor.fetchall()

for fila in datos:
    print(fila)

conn.close()
""", language="python")

st.markdown("""
### ¿Para qué sirve?

Sirve para consultar información guardada en una base SQLite.

En tus proyectos lo usas para:

- NEXUS
- DevData Academy
- Consultas de diagnóstico
""")

st.divider()

st.markdown("## 🌐 Streamlit con Python")

st.code("""
import streamlit as st

st.title("Mi App")
st.write("Hola mundo")

nombre = st.text_input("Escribe tu nombre")

if st.button("Saludar"):
    st.success(f"Hola {nombre}")
""", language="python")

st.markdown("""
### ¿Para qué sirve?

Permite convertir scripts de Python en aplicaciones visuales.
""")

st.divider()

st.markdown("## 🧪 Mini práctica por celdas")

st.markdown("### Celda 1: importar librerías")

st.code("""
import pandas as pd
from pathlib import Path
""", language="python")

st.markdown("### Celda 2: crear datos de ejemplo")

st.code("""
datos = {
    "cliente": ["Ana", "Luis", "Carlos"],
    "ciudad": ["Medellín", "Bogotá", "Cali"],
    "cantidad": [2, 5, 3],
    "precio": [10000, 20000, 15000]
}

df = pd.DataFrame(datos)
df
""", language="python")

st.markdown("### Celda 3: crear columna total")

st.code("""
df["total"] = df["cantidad"] * df["precio"]
df
""", language="python")

st.markdown("### Celda 4: agrupar por ciudad")

st.code("""
resumen = df.groupby("ciudad")["total"].sum().reset_index()
resumen
""", language="python")

st.markdown("### Celda 5: exportar a Excel")

st.code("""
df.to_excel("ventas_practica.xlsx", index=False)
""", language="python")

st.success("""
Esta práctica te ayuda a entender el flujo básico:

crear datos → procesar → calcular → resumir → exportar.
""")

st.divider()

st.markdown("## ⚠️ Errores comunes")

st.warning("""
- No activar el entorno virtual.
- No instalar librerías.
- Escribir mal la ruta del archivo.
- Usar columnas que no existen.
- No revisar `df.columns`.
- No validar tipos de datos.
- No cerrar conexiones SQLite.
- No guardar el archivo antes de ejecutarlo.
""")

st.divider()

st.markdown("## 🧠 Preguntas que debes hacerte antes de crear un script")

st.info("""
1. ¿Qué archivo voy a leer?
2. ¿Qué columnas necesito?
3. ¿Qué datos están sucios?
4. ¿Qué transformación debo hacer?
5. ¿Qué resultado necesito?
6. ¿Dónde voy a guardar la salida?
7. ¿Cómo valido que quedó bien?
""")

st.divider()

st.markdown("## 📌 Checklist para desarrollar en Python")

st.success("""
1. Crear carpeta del proyecto.
2. Crear entorno virtual.
3. Instalar librerías.
4. Crear script.
5. Definir rutas.
6. Cargar datos.
7. Revisar columnas.
8. Limpiar datos.
9. Transformar datos.
10. Exportar resultados.
11. Probar.
12. Documentar.
13. Subir a GitHub.
""")