import streamlit as st

st.set_page_config(page_title="Excel", page_icon="📘", layout="wide")

st.title("📘 Excel")
st.subheader("Tablas estructuradas, fórmulas, Power Query, macros y exportaciones")

st.markdown("""
## 📘 ¿Qué es Excel?

Excel es una herramienta clave para organizar, validar, transformar y analizar datos.

En tus desarrollos lo has usado para:

- Exportar datos desde Python
- Crear tablas estructuradas
- Preparar archivos para Power BI
- Usar fórmulas de validación
- Automatizar tareas con macros
- Revisar inventarios, ANS, reportes y controles
""")

st.divider()

st.markdown("## 🧠 Cómo pensar Excel en proyectos reales")

st.info("""
Excel no debe verse solo como una hoja manual.

En proyectos empresariales puede funcionar como:

1. Entrada de datos
2. Archivo intermedio
3. Archivo de validación
4. Salida de reportes
5. Fuente para Power BI
""")

st.divider()

st.markdown("## 📊 Tablas estructuradas")

st.markdown("""
Una tabla estructurada permite que Excel reconozca los datos como una entidad formal.

Ventajas:

- Los filtros quedan activos
- Las fórmulas se expanden automáticamente
- Power BI detecta mejor los datos
- Power Query trabaja más ordenado
- El archivo queda más profesional
""")

st.code("""
Atajo:
CTRL + T

Recomendación:
Marcar "La tabla tiene encabezados"
""", language="text")

st.markdown("### Buenas prácticas")

st.success("""
- Usar encabezados claros
- No combinar celdas
- No dejar filas vacías dentro de la tabla
- No usar títulos encima de los datos
- Evitar colores manuales como única lógica
- Nombrar tablas de forma clara: TablaVentas, TablaClientes, TablaMovimientos
""")

st.divider()

st.markdown("## 🐍 Exportar Excel desde Python con pandas")

st.code("""
import pandas as pd

df.to_excel("reporte.xlsx", index=False)
""", language="python")

st.markdown("""
### Explicación

- `to_excel()` exporta un DataFrame a Excel.
- `index=False` evita exportar la columna índice de pandas.
""")

st.markdown("### Exportar varias hojas")

st.code("""
with pd.ExcelWriter("reporte.xlsx", engine="openpyxl") as writer:
    ventas.to_excel(writer, sheet_name="Ventas", index=False)
    clientes.to_excel(writer, sheet_name="Clientes", index=False)
    productos.to_excel(writer, sheet_name="Productos", index=False)
""", language="python")

st.info("""
Esto es útil cuando quieres entregar un archivo con varias tablas relacionadas.
""")

st.divider()

st.markdown("## 📋 Convertir exportación en tabla estructurada con openpyxl")

st.code("""
from openpyxl import load_workbook
from openpyxl.worksheet.table import Table, TableStyleInfo

archivo = "reporte.xlsx"

wb = load_workbook(archivo)
ws = wb["Ventas"]

tabla = Table(displayName="TablaVentas", ref=ws.dimensions)

estilo = TableStyleInfo(
    name="TableStyleMedium9",
    showFirstColumn=False,
    showLastColumn=False,
    showRowStripes=True,
    showColumnStripes=False
)

tabla.tableStyleInfo = estilo
ws.add_table(tabla)

wb.save(archivo)
""", language="python")

st.markdown("""
### ¿Para qué sirve?

Convierte el rango exportado desde Python en una tabla formal de Excel.
""")

st.warning("""
Cada tabla debe tener un nombre único. No puedes repetir `TablaVentas` en varias hojas.
""")

st.divider()

st.markdown("## 🧹 Power Query")

st.markdown("""
Power Query sirve para limpiar y transformar datos dentro de Excel sin modificar manualmente el origen.

Lo puedes usar para:

- Quitar columnas
- Cambiar tipos de datos
- Reemplazar valores
- Quitar duplicados
- Unir consultas
- Combinar archivos de una carpeta
- Cargar como solo conexión
- Enviar al modelo de datos
""")

st.markdown("### Tips importantes")

st.success("""
- Cargar tablas auxiliares como "solo conexión".
- Usar nombres claros para consultas.
- Revisar pasos aplicados.
- No depender de rutas rotas.
- Validar tipos de datos antes de cargar.
""")

st.divider()

st.markdown("## 🧮 Fórmulas esenciales")

st.markdown("### SUMA")

st.code("""
=SUMA(A2:A100)
""", language="text")

st.markdown("### PROMEDIO")

st.code("""
=PROMEDIO(B2:B100)
""", language="text")

st.markdown("### SI")

st.code("""
=SI(C2>100000;"Alto";"Bajo")
""", language="text")

st.markdown("### SI.ERROR")

st.code("""
=SI.ERROR(A2/B2;0)
""", language="text")

st.markdown("### CONTAR.SI")

st.code("""
=CONTAR.SI(A2:A100;"Cumplido")
""", language="text")

st.markdown("### SUMAR.SI.CONJUNTO")

st.code("""
=SUMAR.SI.CONJUNTO(Ventas[Total];Ventas[Ciudad];"Medellín";Ventas[Año];2025)
""", language="text")

st.markdown("### CONTAR.SI.CONJUNTO")

st.code("""
=CONTAR.SI.CONJUNTO(Ventas[Ciudad];"Medellín";Ventas[Estado];"Cumplido")
""", language="text")

st.markdown("### BUSCARX")

st.code("""
=BUSCARX(A2;Productos[Codigo];Productos[Nombre];"No encontrado")
""", language="text")

st.markdown("### TEXTO")

st.code("""
=TEXTO(A2;"aaaa-mm")
""", language="text")

st.markdown("### DERECHA / IZQUIERDA / EXTRAE")

st.code("""
=IZQUIERDA(A2;4)
=DERECHA(A2;3)
=EXTRAE(A2;2;5)
""", language="text")

st.markdown("### ESPACIOS / LIMPIAR")

st.code("""
=ESPACIOS(A2)
=LIMPIAR(A2)
""", language="text")

st.markdown("### MAYUSC / MINUSC / NOMPROPIO")

st.code("""
=MAYUSC(A2)
=MINUSC(A2)
=NOMPROPIO(A2)
""", language="text")

st.divider()

st.markdown("## 🚦 Fórmulas útiles para control operativo")

st.markdown("### Estado por cumplimiento")

st.code("""
=SI([@Dias_Restantes]<0;"VENCIDO";SI([@Dias_Restantes]<=2;"ALERTA";"A TIEMPO"))
""", language="text")

st.markdown("### Validar vacío")

st.code("""
=SI([@Tecnico]="";"SIN TÉCNICO";"OK")
""", language="text")

st.markdown("### Diferencia entre cantidades")

st.code("""
=[@Cantidad_Fenix]-[@Cantidad_Elite]
""", language="text")

st.markdown("### Clasificar diferencia")

st.code("""
=SI([@Diferencia]=0;"OK";SI([@Diferencia]>0;"FALTANTE EN ELITE";"EXCESO EN ELITE"))
""", language="text")

st.markdown("### Concatenar llave")

st.code("""
=[@Pedido]&"-"&[@Codigo]
""", language="text")

st.info("""
Las llaves concatenadas sirven para cruzar datos entre fuentes como FÉNIX y planillas internas.
""")

st.divider()

st.markdown("## 🤖 Macros básicas")

st.markdown("""
Las macros permiten automatizar tareas repetitivas en Excel.

Útiles para:

- Actualizar datos
- Limpiar rangos
- Exportar reportes
- Aplicar formatos
- Crear botones
""")

st.markdown("### Macro simple")

st.code("""
Sub Saludar()
    MsgBox "Proceso ejecutado correctamente"
End Sub
""", language="vbnet")

st.markdown("### Limpiar rango")

st.code("""
Sub LimpiarDatos()
    Sheets("Datos").Range("A2:Z1000").ClearContents
End Sub
""", language="vbnet")

st.markdown("### Actualizar todas las conexiones")

st.code("""
Sub ActualizarTodo()
    ThisWorkbook.RefreshAll
    MsgBox "Datos actualizados correctamente"
End Sub
""", language="vbnet")

st.markdown("### Macro para detectar color y escribir zona")

st.code("""
Sub ActualizarZonas()

    Dim celda As Range
    Dim color As Long

    For Each celda In Range("B2:B5000")

        color = celda.Interior.Color

        Select Case color

            Case 5296274
                celda.Offset(0, -1).Value = "AJIZAL (Itagui)"

            Case 3305961
                celda.Offset(0, -1).Value = "BARRIO NUEVO (Medellin)"

            Case Else
                celda.Offset(0, -1).Value = "SIN CLASIFICAR"

        End Select

    Next celda

    MsgBox "Zonas actualizadas correctamente"

End Sub
""", language="vbnet")

st.warning("""
Las macros no viajan igual que un acceso directo. Si compartes el archivo, el usuario debe tener habilitadas macros y rutas correctas en su equipo.
""")

st.divider()

st.markdown("## 📦 Excel como salida para Power BI")

st.markdown("""
Cuando Excel será fuente para Power BI:

- Usa tablas estructuradas.
- Evita celdas combinadas.
- Usa nombres de columnas limpios.
- No dejes totales manuales dentro de la tabla.
- Evita comentarios dentro del rango.
- Mantén tipos de datos consistentes.
""")

st.markdown("### Columnas recomendadas")

st.code("""
id
fecha
codigo
descripcion
cantidad
precio
total
estado
ciudad
responsable
""", language="text")

st.divider()

st.markdown("## ⚠️ Errores comunes")

st.warning("""
- No convertir rango en tabla.
- Cambiar nombres de columnas sin actualizar Power Query.
- Tener columnas con tipos mezclados.
- Dejar filas vacías.
- Usar fórmulas arrastradas manualmente.
- Guardar archivo con macros como .xlsx en vez de .xlsm.
- No validar resultados antes de entregar.
""")

st.divider()

st.markdown("## ✅ Checklist Excel profesional")

st.success("""
1. Datos en tabla estructurada.
2. Encabezados claros.
3. Sin filas vacías.
4. Sin celdas combinadas.
5. Tipos de datos correctos.
6. Fórmulas validadas.
7. Archivo guardado en formato correcto.
8. Si usa macros, guardar como .xlsm.
9. Si va a Power BI, validar nombres de columnas.
10. Hacer backup antes de cambios grandes.
""")